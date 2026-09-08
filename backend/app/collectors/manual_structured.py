from __future__ import annotations

import csv
from dataclasses import dataclass
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

IMF_WEO_COUNTRIES = ("COD", "ZWE", "ZMB", "ZAF")
IMF_WEO_YEARS = tuple(range(2015, 2025))
IMF_WEO_DATASET = "IMF.RES:WEO(9.0.0)"
IMF_WEO_INDICATOR = "GGXWDG_NGDP"
IMF_WEO_METRIC = "general_government_gross_debt_percent_gdp"

NBS_YEARS = tuple(range(2016, 2026))
NBS_INDICATORS = (
    ("工业增加值 (亿元)", "industrial_value_added", "亿元"),
    ("十种有色金属产量 (万吨)", "ten_nonferrous_metals_output", "万吨"),
    ("精炼铜产量 (万吨)", "refined_copper_output", "万吨"),
    ("原铝 (电解铝) 产量 (万吨)", "primary_aluminium_output", "万吨"),
    (
        "采矿业固定资产投资 (不含农户) 比上年增长 (%)",
        "mining_fixed_asset_investment_yoy",
        "percent",
    ),
    (
        "有色金属冶炼和压延加工业固定资产投资 (不含农户) 比上年增长 (%)",
        "nonferrous_smelting_rolling_fixed_asset_investment_yoy",
        "percent",
    ),
)

MOFCOM_TRADE_MONTHS = tuple(year * 100 + month for year in range(2020, 2025) for month in range(1, 13))
MOFCOM_TRADE_METRICS = (
    ("trade_total_amount", "hundred_million_USD"),
    ("trade_total_yoy", "percent"),
    ("export_amount", "hundred_million_USD"),
    ("export_yoy", "percent"),
    ("import_amount", "hundred_million_USD"),
    ("import_yoy", "percent"),
)
MOFCOM_ODI_MONTHS = tuple(year * 100 + month for year in range(2015, 2025) for month in range(1, 13))
MOFCOM_ODI_METRICS = (
    ("country_region_count", "count"),
    ("enterprise_count", "count"),
    ("investment_amount", "hundred_million_USD"),
    ("investment_yoy", "percent"),
)


class ManualStructuredParseError(ValueError):
    def __init__(self, code: str, message: str) -> None:
        super().__init__(message)
        self.code = code
        self.message = message


@dataclass(frozen=True, slots=True)
class ManualStructuredCell:
    country_iso3: str
    period: int
    indicator_code: str
    metric_code: str
    value: Decimal | None
    unit: str
    source_status: str
    missing_reason: str | None
    quality_flags: dict[str, Any]
    partner_iso3: str | None = None
    currency: str | None = None
    price_basis: str | None = None
    is_aggregate: bool | None = None


@dataclass(frozen=True, slots=True)
class ManualStructuredParseResult:
    provider_version: str
    source_release_date: date | None
    expected_count: int
    returned_count: int
    valued_count: int
    source_null_count: int
    not_returned_count: int
    source_metadata: dict[str, Any]
    observations: tuple[ManualStructuredCell, ...]


def parse_imf_weo_csv(path: str | Path) -> ManualStructuredParseResult:
    csv_path = Path(path)
    try:
        handle = csv_path.open("r", encoding="utf-8-sig", newline="")
    except OSError as exc:
        raise ManualStructuredParseError("file_unreadable", "IMF WEO CSV cannot be read") from exc

    with handle:
        reader = csv.DictReader(handle)
        required = {
            "DATASET",
            "SERIES_CODE",
            "OBS_MEASURE",
            "COUNTRY",
            "INDICATOR",
            "FREQUENCY",
            "SCALE",
            "UNIT",
            "PUBLICATION_DATE",
            "UPDATE_DATE",
            "LATEST_ACTUAL_ANNUAL_DATA",
            *{str(year) for year in IMF_WEO_YEARS},
        }
        if reader.fieldnames is None or not required.issubset(reader.fieldnames):
            raise ManualStructuredParseError(
                "schema_mismatch",
                "IMF WEO CSV does not expose the reviewed columns",
            )

        expected_series = {f"{country}.{IMF_WEO_INDICATOR}.A": country for country in IMF_WEO_COUNTRIES}
        rows: dict[str, dict[str, str]] = {}
        for row in reader:
            series_code = (row.get("SERIES_CODE") or "").strip()
            if series_code not in expected_series:
                continue
            if series_code in rows:
                raise ManualStructuredParseError(
                    "duplicate_series",
                    f"IMF WEO CSV repeats {series_code}",
                )
            rows[series_code] = row

    if set(rows) != set(expected_series):
        missing = sorted(set(expected_series) - set(rows))
        raise ManualStructuredParseError(
            "incomplete_scope",
            f"IMF WEO CSV is missing reviewed series: {missing}",
        )

    publication_dates: set[datetime] = set()
    update_dates: set[datetime] = set()
    observations: list[ManualStructuredCell] = []
    latest_actual: dict[str, str] = {}
    country_labels: dict[str, str] = {}
    for series_code, country_iso3 in expected_series.items():
        row = rows[series_code]
        expected_fields = {
            "DATASET": IMF_WEO_DATASET,
            "OBS_MEASURE": "OBS_VALUE",
            "INDICATOR": "Gross debt, General government, Percent of GDP",
            "FREQUENCY": "Annual",
            "SCALE": "Units",
            "UNIT": "Percent",
        }
        for field, expected in expected_fields.items():
            if (row.get(field) or "").strip() != expected:
                raise ManualStructuredParseError(
                    "scope_mismatch",
                    f"IMF WEO {series_code} has unexpected {field}",
                )
        publication_dates.add(_parse_iso_datetime(row["PUBLICATION_DATE"], "PUBLICATION_DATE"))
        update_dates.add(_parse_iso_datetime(row["UPDATE_DATE"], "UPDATE_DATE"))
        latest_actual[country_iso3] = (row.get("LATEST_ACTUAL_ANNUAL_DATA") or "").strip()
        country_labels[country_iso3] = (row.get("COUNTRY") or "").strip()
        for year in IMF_WEO_YEARS:
            value = _decimal(row[str(year)], f"{series_code}/{year}", allow_blank=False)
            observations.append(
                ManualStructuredCell(
                    country_iso3=country_iso3,
                    period=year,
                    indicator_code=IMF_WEO_INDICATOR,
                    metric_code=IMF_WEO_METRIC,
                    value=value,
                    unit="percent_of_GDP",
                    source_status="normal",
                    missing_reason=None,
                    quality_flags={
                        "series_code": series_code,
                        "country_label": country_labels[country_iso3],
                        "latest_actual_annual_data": latest_actual[country_iso3],
                    },
                )
            )

    if len(publication_dates) != 1 or len(update_dates) != 1:
        raise ManualStructuredParseError(
            "release_mismatch",
            "IMF WEO target series do not share one publication and update date",
        )
    publication_date = next(iter(publication_dates))
    update_date = next(iter(update_dates))
    return ManualStructuredParseResult(
        provider_version=IMF_WEO_DATASET,
        source_release_date=publication_date.date(),
        expected_count=40,
        returned_count=40,
        valued_count=40,
        source_null_count=0,
        not_returned_count=0,
        source_metadata={
            "dataset": IMF_WEO_DATASET,
            "indicator": IMF_WEO_INDICATOR,
            "publication_date": publication_date.isoformat(),
            "update_date": update_date.isoformat(),
            "country_labels": country_labels,
            "latest_actual_annual_data": latest_actual,
        },
        observations=tuple(observations),
    )


def parse_nbs_annual_xlsx(path: str | Path) -> ManualStructuredParseResult:
    xlsx_path = Path(path)
    try:
        workbook = load_workbook(xlsx_path, read_only=True, data_only=True)
    except (OSError, ValueError) as exc:
        raise ManualStructuredParseError(
            "file_unreadable",
            "国家统计局年度数据 Excel cannot be read",
        ) from exc
    try:
        if workbook.sheetnames != ["Sheet1"]:
            raise ManualStructuredParseError(
                "schema_mismatch",
                "国家统计局年度数据 Excel must contain only Sheet1",
            )
        sheet = workbook["Sheet1"]
        if sheet.max_row != 12 or sheet.max_column != 11:
            raise ManualStructuredParseError(
                "schema_mismatch",
                "国家统计局年度数据 Excel dimensions differ from the reviewed export",
            )
        if sheet["A1"].value != "数据库：年度数据" or sheet["A2"].value != "时间：2016年-2025年":
            raise ManualStructuredParseError(
                "scope_mismatch",
                "国家统计局年度数据 Excel title or period differs from the reviewed export",
            )
        expected_headers = ["指标", *[f"{year}年" for year in reversed(NBS_YEARS)]]
        actual_headers = [sheet.cell(row=3, column=column).value for column in range(1, 12)]
        if actual_headers != expected_headers:
            raise ManualStructuredParseError(
                "schema_mismatch",
                "国家统计局年度数据 Excel year headers differ from 2016-2025",
            )

        observations: list[ManualStructuredCell] = []
        source_null_count = 0
        for row_index, (label, indicator_code, unit) in enumerate(NBS_INDICATORS, start=4):
            source_label = str(sheet.cell(row=row_index, column=1).value or "").strip()
            if source_label != label:
                raise ManualStructuredParseError(
                    "scope_mismatch",
                    f"国家统计局年度数据 Excel row {row_index} indicator differs from the reviewed scope",
                )
            for column_index, year in enumerate(reversed(NBS_YEARS), start=2):
                raw_value = sheet.cell(row=row_index, column=column_index).value
                value = _decimal(raw_value, f"{label}/{year}", allow_blank=True)
                is_null = value is None
                source_null_count += int(is_null)
                observations.append(
                    ManualStructuredCell(
                        country_iso3="CHN",
                        period=year,
                        indicator_code=indicator_code,
                        metric_code=indicator_code,
                        value=value,
                        unit=unit,
                        source_status="source_null" if is_null else "normal",
                        missing_reason="source_null" if is_null else None,
                        quality_flags={
                            "source_indicator_label": source_label,
                            "source_indicator_id": None,
                            "geography": "全国",
                        },
                    )
                )

        notes = [
            str(sheet.cell(row=row_index, column=1).value).strip()
            for row_index in range(10, 13)
            if sheet.cell(row=row_index, column=1).value
        ]
    finally:
        workbook.close()

    return ManualStructuredParseResult(
        provider_version="国家统计局全国年度数据-2016-2025",
        source_release_date=None,
        expected_count=60,
        returned_count=60,
        valued_count=60 - source_null_count,
        source_null_count=source_null_count,
        not_returned_count=0,
        source_metadata={
            "database": "年度数据",
            "geography": "中国/全国",
            "years": [2016, 2025],
            "indicator_count": 6,
            "source_indicator_ids_available": False,
            "notes": notes,
        },
        observations=tuple(observations),
    )


def parse_mofcom_south_africa_trade_csv(path: str | Path) -> ManualStructuredParseResult:
    rows = _read_trimmed_csv(path, "商务部南非货物进出口 CSV")
    expected_header = [
        ["南非历年货物进出口分国别（地区）统计"],
        ["日期", "进出口（亿美元）", "", "出口（亿美元）", "", "进口（亿美元）"],
        ["", "累计金额", "同比（%）", "累计金额", "同比（%）", "累计金额", "同比（%）"],
        [],
    ]
    if rows[:4] != expected_header:
        raise ManualStructuredParseError(
            "schema_mismatch",
            "商务部南非货物进出口 CSV headers differ from the reviewed export",
        )

    data_rows = rows[4:]
    if len(data_rows) != len(MOFCOM_TRADE_MONTHS):
        raise ManualStructuredParseError(
            "incomplete_scope",
            "商务部南非货物进出口 CSV must contain 60 monthly rows",
        )

    by_period: dict[int, list[str]] = {}
    for row_number, row in enumerate(data_rows, start=5):
        if len(row) != 7:
            raise ManualStructuredParseError(
                "schema_mismatch",
                f"商务部南非货物进出口 CSV row {row_number} must contain 7 fields",
            )
        period = _parse_chinese_month(row[0], cumulative_prefix=False)
        if period in by_period:
            raise ManualStructuredParseError(
                "duplicate_period",
                f"商务部南非货物进出口 CSV repeats {period}",
            )
        by_period[period] = row
    if set(by_period) != set(MOFCOM_TRADE_MONTHS):
        raise ManualStructuredParseError(
            "incomplete_scope",
            "商务部南非货物进出口 CSV month coverage must be 2020-01 through 2024-12",
        )

    observations: list[ManualStructuredCell] = []
    for period in MOFCOM_TRADE_MONTHS:
        row = by_period[period]
        for column, (metric_code, unit) in enumerate(MOFCOM_TRADE_METRICS, start=1):
            observations.append(
                ManualStructuredCell(
                    country_iso3="CHN",
                    partner_iso3="ZAF",
                    period=period,
                    indicator_code="MOFCOM.GOODS_TRADE_BY_COUNTRY",
                    metric_code=metric_code,
                    value=_decimal(row[column], f"{row[0]}/{metric_code}", allow_blank=False),
                    unit=unit,
                    currency="USD" if metric_code.endswith("_amount") else None,
                    price_basis="current" if metric_code.endswith("_amount") else None,
                    source_status="normal",
                    missing_reason=None,
                    is_aggregate=True,
                    quality_flags={
                        "source_period_label": row[0],
                        "reporting_country": "中国",
                        "partner_country": "南非",
                        "aggregation": "year_to_date",
                        "source_organization": "海关总署",
                    },
                )
            )

    return ManualStructuredParseResult(
        provider_version="商务部商务数据中心-南非货物进出口-2020-2024",
        source_release_date=None,
        expected_count=360,
        returned_count=360,
        valued_count=360,
        source_null_count=0,
        not_returned_count=0,
        source_metadata={
            "table": "南非历年货物进出口分国别（地区）统计",
            "reporting_country": "CHN",
            "partner_country": "ZAF",
            "period_start": 202001,
            "period_end": 202412,
            "frequency": "monthly",
            "aggregation": "year_to_date",
            "source_organization": "海关总署",
        },
        observations=tuple(observations),
    )


def parse_mofcom_nonfinancial_odi_csv(path: str | Path) -> ManualStructuredParseResult:
    rows = _read_trimmed_csv(path, "商务部非金融类对外直接投资 CSV")
    expected_header = [
        ["非金融类对外直接投资统计"],
        ["时间", "国家/地区(个)", "企业数(家)", "金额(亿美元)", "同比(%)"],
        [],
    ]
    if rows[:3] != expected_header:
        raise ManualStructuredParseError(
            "schema_mismatch",
            "商务部非金融类对外直接投资 CSV headers differ from the reviewed export",
        )

    data_rows = rows[3:]
    if len(data_rows) != len(MOFCOM_ODI_MONTHS):
        raise ManualStructuredParseError(
            "incomplete_scope",
            "商务部非金融类对外直接投资 CSV must contain 120 monthly rows",
        )

    by_period: dict[int, list[str]] = {}
    for row_number, row in enumerate(data_rows, start=4):
        if len(row) != 5:
            raise ManualStructuredParseError(
                "schema_mismatch",
                f"商务部非金融类对外直接投资 CSV row {row_number} must contain 5 fields",
            )
        period = _parse_chinese_month(row[0], cumulative_prefix=True)
        if period in by_period:
            raise ManualStructuredParseError(
                "duplicate_period",
                f"商务部非金融类对外直接投资 CSV repeats {period}",
            )
        by_period[period] = row
    if set(by_period) != set(MOFCOM_ODI_MONTHS):
        raise ManualStructuredParseError(
            "incomplete_scope",
            "商务部非金融类对外直接投资 CSV month coverage must be 2015-01 through 2024-12",
        )

    observations: list[ManualStructuredCell] = []
    for period in MOFCOM_ODI_MONTHS:
        row = by_period[period]
        for column, (metric_code, unit) in enumerate(MOFCOM_ODI_METRICS, start=1):
            observations.append(
                ManualStructuredCell(
                    country_iso3="CHN",
                    partner_iso3="WLD",
                    period=period,
                    indicator_code="MOFCOM.NONFINANCIAL_ODI",
                    metric_code=metric_code,
                    value=_decimal(row[column], f"{row[0]}/{metric_code}", allow_blank=False),
                    unit=unit,
                    currency="USD" if metric_code == "investment_amount" else None,
                    price_basis="current" if metric_code == "investment_amount" else None,
                    source_status="normal",
                    missing_reason=None,
                    is_aggregate=True,
                    quality_flags={
                        "source_period_label": row[0],
                        "reporting_country": "中国",
                        "partner_scope": "全球合计",
                        "aggregation": "year_to_date",
                        "source_organization": "商务部合作司",
                    },
                )
            )

    return ManualStructuredParseResult(
        provider_version="商务部商务数据中心-非金融类对外直接投资-2015-2024",
        source_release_date=None,
        expected_count=480,
        returned_count=480,
        valued_count=480,
        source_null_count=0,
        not_returned_count=0,
        source_metadata={
            "table": "非金融类对外直接投资统计",
            "reporting_country": "CHN",
            "partner_scope": "WLD",
            "period_start": 201501,
            "period_end": 202412,
            "frequency": "monthly",
            "aggregation": "year_to_date",
            "source_organization": "商务部合作司",
        },
        observations=tuple(observations),
    )


def _read_trimmed_csv(path: str | Path, label: str) -> list[list[str]]:
    try:
        handle = Path(path).open("r", encoding="utf-8-sig", newline="")
    except OSError as exc:
        raise ManualStructuredParseError("file_unreadable", f"{label} cannot be read") from exc
    with handle:
        rows = []
        for raw_row in csv.reader(handle):
            row = [cell.strip() for cell in raw_row]
            while row and not row[-1]:
                row.pop()
            rows.append(row)
    return rows


def _parse_chinese_month(value: str, *, cumulative_prefix: bool) -> int:
    suffix = value.strip()
    if "年" not in suffix or not suffix.endswith("月"):
        raise ManualStructuredParseError("invalid_period", f"invalid monthly period: {value}")
    year_text, month_text = suffix[:-1].split("年", maxsplit=1)
    try:
        if cumulative_prefix:
            if month_text == "1":
                month_number = 1
            elif month_text.startswith("1-"):
                month_number = int(month_text[2:])
            else:
                raise ManualStructuredParseError(
                    "invalid_period",
                    f"invalid cumulative period: {value}",
                )
        else:
            month_number = int(month_text)
        year = int(year_text)
    except ValueError as exc:
        raise ManualStructuredParseError(
            "invalid_period",
            f"invalid monthly period: {value}",
        ) from exc
    if not 1 <= month_number <= 12:
        raise ManualStructuredParseError("invalid_period", f"invalid monthly period: {value}")
    return year * 100 + month_number


def _parse_iso_datetime(value: str, field: str) -> datetime:
    try:
        return datetime.fromisoformat(value.replace("Z", "+00:00"))
    except (AttributeError, ValueError) as exc:
        raise ManualStructuredParseError(
            "invalid_release_date",
            f"IMF WEO {field} is invalid",
        ) from exc


def _decimal(value: object, label: str, *, allow_blank: bool) -> Decimal | None:
    if value is None or (isinstance(value, str) and not value.strip()):
        if allow_blank:
            return None
        raise ManualStructuredParseError("missing_value", f"{label} is blank")
    if isinstance(value, bool):
        raise ManualStructuredParseError("invalid_value", f"{label} is not numeric")
    try:
        parsed = Decimal(str(value).strip())
    except (InvalidOperation, ValueError) as exc:
        raise ManualStructuredParseError("invalid_value", f"{label} is not numeric") from exc
    if not parsed.is_finite():
        raise ManualStructuredParseError("invalid_value", f"{label} is not finite")
    return parsed
