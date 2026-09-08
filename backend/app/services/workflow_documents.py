"""Append-only readable S03/S05 drafts, source attachments and immutable confirmations."""

import copy
import csv
import io
import re
from collections import defaultdict
from datetime import UTC, datetime

from sqlalchemy import select

from app.models import CapabilityRun, CapabilityRunRevision, FieldMaterial, ResearchCase
from app.models.skill_workflows import FieldAnnotation, FieldMaterialProfile, FieldSegment

DOCUMENT_KIND = "workflow_document"
CITATION = re.compile(r"\[E(\d+)\]")


def revisions(db, run_id):
    return [
        row
        for row in db.scalars(
            select(CapabilityRunRevision)
            .where(CapabilityRunRevision.capability_run_id == run_id)
            .order_by(CapabilityRunRevision.revision_no)
        )
        if (row.artifact or {}).get("kind") == DOCUMENT_KIND
    ]


def latest(db, run_id):
    return next((r for r in reversed(revisions(db, run_id)) if r.artifact["state"] != "proposal"), None)


def append(db, run, actor_id, document, state, note, *, parent_id=None):
    db.scalar(select(CapabilityRun).where(CapabilityRun.id == run.id).with_for_update())
    prior = db.scalar(
        select(CapabilityRunRevision)
        .where(CapabilityRunRevision.capability_run_id == run.id)
        .order_by(CapabilityRunRevision.revision_no.desc())
        .limit(1)
    )
    row = CapabilityRunRevision(
        capability_run_id=run.id,
        revision_no=prior.revision_no + 1 if prior else 1,
        artifact={
            "kind": DOCUMENT_KIND,
            "state": state,
            "parent_id": parent_id,
            "document": copy.deepcopy(document),
        },
        note=note,
        created_by=actor_id,
    )
    db.add(row)
    db.flush()
    return row


def source_evidence(db, run):
    evidence, materials = [], []
    for mid in (run.output or {}).get("material_ids", []):
        material = db.get(FieldMaterial, mid)
        profile = db.get(FieldMaterialProfile, mid)
        context = profile.context if profile else {}
        materials.append(
            {
                "material_id": mid,
                "research_case_id": run.research_case_id,
                "title": material.title,
                "captured_on": str(material.captured_on or "未补充"),
                "context": context,
                "content_type": material.content_type,
            }
        )
    for segment in db.scalars(
        select(FieldSegment).where(FieldSegment.run_id == run.id).order_by(FieldSegment.id)
    ):
        annotation = db.scalar(
            select(FieldAnnotation)
            .where(FieldAnnotation.segment_id == segment.id)
            .order_by(FieldAnnotation.revision_no.desc())
            .limit(1)
        )
        if not annotation or annotation.status == "rejected":
            continue
        evidence.append(
            {
                "id": segment.id,
                "material_id": segment.material_id,
                "raw_version_id": segment.raw_version_id,
                "clean_version_id": segment.clean_version_id,
                "quote": segment.original_text,
                "positions": segment.positions,
                "annotation_id": annotation.id,
                "status": annotation.status,
                **annotation.payload,
            }
        )
    return evidence, materials


def initial_document(db, run):
    case = db.get(ResearchCase, run.research_case_id)
    if run.output.get("type") == "field_research_workflow":
        evidence, materials = source_evidence(db, run)
        themes = defaultdict(list)
        for item in evidence:
            themes[(item.get("themes") or ["待归纳"])[0]].append(item)
        title = f"{case.title} · 田野研究备忘录"
        parts = [
            f"# {title}",
            "## 研究问题与材料范围",
            case.research_question or "研究问题待补充。",
            f"本稿依据所选 {len(materials)} 份资料、{len(evidence)} 条可追溯片段形成。"
            "以下为研究候选发现，材料中的陈述尚不等于已核实事实。",
        ]
        for material in materials:
            c = material["context"]
            parts.append(
                f"《{material['title']}》：田野时间 {material['captured_on']}；"
                f"地点 {c.get('location') or '待补充'}；对象代号 {c.get('participant_alias') or '待补充'}。"
            )
        parts += [
            "## 方法、样本来源与整理方式",
            "原始文件与原文单独保留。文本按完整材料分批整理，候选发现与原文位置逐条对应；"
            "人工批注、整理稿修订和采用状态在证据附件中保留。",
        ]
        parts.extend(
            m["context"].get("sample_note") or f"《{m['title']}》的样本来源尚未补充。" for m in materials
        )
        parts.append("## 主题发现及支持证据")
        for theme, items in themes.items():
            parts.append(f"### {theme}")
            parts.extend(
                f"{item.get('viewpoint') or '该片段的研究意义待研究者说明。'} [E{item['id']}]"
                for item in items
            )
        if not themes:
            parts.append("当前没有可据以形成发现的文本片段；图片或扫描件需补充说明或转写。")
        parts.append("## 分歧、反例与待核实问题")
        conflicts = [x for x in evidence if x.get("verification") == "conflict"]
        parts.extend(f"{x.get('note') or x.get('viewpoint')} [E{x['id']}]" for x in conflicts)
        if not conflicts:
            parts.append(
                "尚未单独标注分歧或反例；这不表示不同材料之间不存在分歧。需继续核对样本差异与陈述来源。"
            )
        parts += ["## 研究者反思、局限与下一步"]
        parts.extend(
            m["context"]["researcher_reflection"]
            for m in materials
            if m["context"].get("researcher_reflection")
        )
        parts.extend(run.output.get("limitations", []))
        parts.append("补充缺失背景，核实关键数字与否定表达，并逐条复核证据后确认正文。")
    else:
        title = f"{case.title} · 专题更新简报"
        evidence, materials = [], []
        parts = [
            f"# {title}",
            "## 本轮范围",
            case.research_question or "追踪问题待补充。",
            "## 新增与实质变化",
        ]
        for index, observation in enumerate(run.output.get("observations", []), 1):
            snapshot = observation["snapshot"]
            if observation.get("observation") not in {"new", "updated"}:
                continue
            evidence.append(
                {
                    "id": index,
                    "observation_id": observation.get("observation_id"),
                    "item_id": observation["item_id"],
                    "snapshot": snapshot,
                    "status": "candidate",
                    "quote": "",
                    "positions": {},
                    "title": snapshot.get("title", ""),
                }
            )
            parts.append(
                f"{snapshot.get('title', '待补充标题')}：来源 {snapshot.get('source_name', '未登记')}；"
                f"实际获取时间 {snapshot.get('observed_at', '未登记')}。 [E{index}]"
            )
        parts += [
            "## 影响与局限",
            "以上材料尚待研究者审核。来源访问失败与没有更新应结合运行记录分别判断。",
            "## 下一步",
            "核对新增材料并判断是否采用；未生成自动采用或正式发布。",
        ]
    return {
        "title": title,
        "markdown": "\n\n".join(str(x) for x in parts if x),
        "evidence": evidence,
        "materials": materials,
        "source_run_id": run.id,
        "workflow_type": run.output.get("type"),
        "created_at": datetime.now(UTC).isoformat(),
    }


def validate_body(document, markdown):
    known = {int(x["id"]) for x in document["evidence"]}
    used = {int(x) for x in CITATION.findall(markdown)}
    if used - known:
        raise ValueError("正文引用了不存在的证据编号")
    # Quoted originals live in a separate immutable attachment, never in the editable body.
    if len(markdown) > 200_000:
        raise ValueError("正文过长")


def apply_tracking_prompt(db, run, document, instruction):
    """Use the scheduled instruction with bounded metadata and the same field AI boundary."""
    from contextlib import ExitStack

    from pydantic import BaseModel, ConfigDict, Field

    from app.core.config import get_settings
    from app.services import field_access, field_ai_access, workflow_runtime

    document["writing_instruction"] = instruction
    if get_settings().agent_runtime == "evidence_only":
        document["generation_mode"] = "metadata_template"
        return document

    class Draft(BaseModel):
        model_config = ConfigDict(extra="forbid")
        markdown: str = Field(min_length=20, max_length=30000)

    metadata = [
        {
            "evidence_id": item["id"],
            **{
                key: item.get("snapshot", {}).get(key)
                for key in ("title", "source_name", "source_url", "published_at", "observed_at", "issues")
            },
        }
        for item in document["evidence"]
    ]
    with ExitStack() as stack:
        for material_id in field_access.run_material_ids(db, run):
            stack.enter_context(field_ai_access.authorized_source(db, run, material_id))
        draft = workflow_runtime.invoke(
            Draft,
            {"instruction": instruction, "outline": document["markdown"], "evidence": metadata},
            "依据研究者提示词整理专题简报。只掌握题名、来源、日期和获取状态，不能从题名推断正文结论。"
            "每项材料使用 [E编号] 引用，不输出原文引句。新增推断必须明确写待核实。"
            "保留实际采集时间与材料发表时间的区别、失败说明和人工审阅边界。不联网、不调用工具。",
        )
    validate_body(document, draft.markdown)
    required = {x["id"] for x in document["evidence"]}
    if {int(x) for x in CITATION.findall(draft.markdown)} != required:
        raise ValueError("专题简报引用未覆盖本轮材料，未保存不完整草稿")
    document["markdown"] = draft.markdown
    document["generation_mode"] = "instruction_applied"
    return document


def markdown_export(document):
    lines = [document["markdown"], "\n## 证据附件"]
    for item in document["evidence"]:
        lines.append(
            f"\n### E{item['id']} · {item.get('title') or '资料 #' + str(item.get('material_id', ''))}"
        )
        if item.get("quote"):
            lines.append("\n".join("> " + line for line in item["quote"].splitlines()))
        lines.append(
            f"原文版本：{item.get('raw_version_id', '—')}；定位：{item.get('positions', {})}；"
            f"审核状态：{item['status']}。"
        )
        if item.get("snapshot"):
            s = item["snapshot"]
            lines.append(
                f"来源：{s.get('source_name', '')}；地址：{s.get('source_url', '')}；"
                f"获取时间：{s.get('observed_at', '')}。"
            )
    return "\n\n".join(lines)


def csv_export(document):
    output = io.StringIO()
    writer = csv.writer(output)
    writer.writerow(["证据编号", "材料/观察", "主题", "对象代号", "原文", "定位", "审核状态"])
    aliases = {
        m["material_id"]: m["context"].get("participant_alias", "") for m in document.get("materials", [])
    }

    def safe(value):
        value = str(value)
        return "'" + value if value.startswith(("=", "+", "-", "@", "\t", "\r")) else value

    for item in document["evidence"]:
        writer.writerow(
            [
                safe(x)
                for x in [
                    f"E{item['id']}",
                    item.get("material_id", item.get("observation_id")),
                    "、".join(item.get("themes", [])),
                    aliases.get(item.get("material_id"), ""),
                    item.get("quote", ""),
                    item.get("positions", {}),
                    item["status"],
                ]
            ]
        )
    return "\ufeff" + output.getvalue()


def docx_export(document):
    from docx import Document
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Cm, Pt, RGBColor

    doc = Document()
    section = doc.sections[0]
    section.top_margin = section.bottom_margin = Cm(2.2)
    section.left_margin = section.right_margin = Cm(2.4)
    section.page_width, section.page_height = Cm(21), Cm(29.7)
    for name in ("Normal", "Title", "Heading 1", "Heading 2", "Heading 3", "Quote", "Footer"):
        style = doc.styles[name]
        style.font.name = "Songti SC"
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.font.italic = False
        fonts = style.element.get_or_add_rPr().rFonts
        for key in list(fonts.attrib):
            del fonts.attrib[key]
        for key in ("ascii", "hAnsi", "eastAsia", "cs"):
            fonts.set(qn("w:" + key), "Songti SC")
        for border in style.element.xpath("./w:pPr/w:pBdr"):
            border.getparent().remove(border)
    normal = doc.styles["Normal"]
    normal.font.size = Pt(11)
    normal.paragraph_format.line_spacing = 1.3
    for line in document["markdown"].splitlines():
        if not line.strip():
            continue
        match = re.match(r"^(#{1,3})\s+(.*)", line)
        if match:
            doc.add_heading(match[2], level=len(match[1]) - 1)
        else:
            doc.add_paragraph(line)
    doc.add_page_break()
    doc.add_heading("证据与研究过程", 1)
    for item in document["evidence"]:
        doc.add_heading(
            f"E{item['id']} · {item.get('title') or '材料 #' + str(item.get('material_id', ''))}", 2
        )
        if item.get("quote"):
            quotation = doc.add_paragraph(item["quote"], "Quote")
            quotation.paragraph_format.keep_with_next = True
        doc.add_paragraph(
            f"原文版本 #{item.get('raw_version_id', '—')}；定位 {item.get('positions', {})}；"
            f"状态 {item['status']}"
        )
        if item.get("snapshot"):
            doc.add_paragraph(str(item["snapshot"].get("source_url", "")))
    footer = section.footer.paragraphs[0]
    footer.alignment = 2
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    footer._p.append(field)
    out = io.BytesIO()
    doc.save(out)
    return out.getvalue()
